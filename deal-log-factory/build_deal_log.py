#!/usr/bin/env python3
"""build_deal_log.py — Gallatin Ford Deal Log generator (SPEC.md is the contract).

Architecture: config-driven TRANSFORM. The committed July master is loaded as the
style/structure donor; every month-dependent artifact (names, dates, serials, pace
numbers, calendar geometry, DOW labels, Sunday treatment) is rewritten from BUILD
CONFIG, then the Gate-1-approved deltas are applied. Future months = edit CONFIG.

Usage:
    python3 build_deal_log.py                # clean build -> OUTPUT_FILE
    python3 build_deal_log.py --test-deals   # + seed the 7-deal Gate-3 matrix on day 1
"""
import argparse
import datetime as dt
import re
import sys
import warnings
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

warnings.filterwarnings("ignore")
HERE = Path(__file__).resolve().parent

# ============================================================ BUILD CONFIG
CONFIG = dict(
    MONTH="August", YEAR=2026, MONTH_NUM=8, ABBR="Aug",
    DAYS=31, FIRST_DOW="Saturday",
    CLOSED_DAYS=[2, 9, 16, 23, 30],
    SELLING_DAYS=26,
    STORE="Gallatin Ford",
    SOURCE_FILE=HERE / "source/July_2026_Ford_Deal_Log_Master.xlsx",
    OUTPUT_FILE=HERE / "August_2026_Ford_Deal_Log_Master.xlsx",
)
DONOR = dict(ABBR="Jul", NAME="July", MONTH_NUM=7, DAYS=31, SELLING=27)

# Tab colors (SPEC §2)
COLOR_WEEKDAY, COLOR_SATURDAY, COLOR_SUNDAY = "0B1F4D", "2E4A8F", "B8C4DC"
YELLOW = "FFF6BE"

# ============================================================ derived month facts
def month_facts(cfg):
    y, m, n = cfg["YEAR"], cfg["MONTH_NUM"], cfg["DAYS"]
    closed = set(cfg["CLOSED_DAYS"])
    days = list(range(1, n + 1))
    selling = [d for d in days if d not in closed]
    assert len(selling) == cfg["SELLING_DAYS"], "SELLING_DAYS inconsistent with CLOSED_DAYS"
    first = dt.date(y, m, 1)
    assert first.strftime("%A") == cfg["FIRST_DOW"], "FIRST_DOW inconsistent"
    assert all(dt.date(y, m, d).strftime("%A") == "Sunday" for d in closed), "CLOSED_DAYS not Sundays"
    base = dt.date(1899, 12, 30)
    return dict(
        days=days, closed=closed, selling=selling,
        elapsed={d: len([s for s in selling if s <= d]) for d in days},
        left={d: len([s for s in selling if s >= d]) for d in days},
        serial={d: (dt.date(y, m, d) - base).days for d in days},
        dow={d: dt.date(y, m, d).strftime("%A") for d in days},
        weeks={d: (d + (7 - (8 - first.isoweekday()) % 7) - 2 + 7) // 7 if False else None for d in days},
    )

def week_index(cfg, d):
    """0-based calendar week row for day d (weeks start Sunday)."""
    first = dt.date(cfg["YEAR"], cfg["MONTH_NUM"], 1)
    dow0 = (first.isoweekday()) % 7          # Sun=0 .. Sat=6
    return (d - 1 + dow0) // 7

def dow_col(cfg, d):
    """Calendar column letter for day d (Sun..Sat -> B,D,F,H,J,L,N)."""
    wd = (dt.date(cfg["YEAR"], cfg["MONTH_NUM"], d).isoweekday()) % 7
    return "BDFHJLN"[wd]

F = month_facts(CONFIG)
A = CONFIG["ABBR"]                      # 'Aug'
TABS = [f"{A} {d}" for d in F["days"]]
LAST = TABS[-1]

def chain(fmt, days=None):
    """31-term '+'-joined chain over day tabs. fmt uses {t} for quoted tab name."""
    return "=" + "+".join(fmt.format(t=f"'{A} {d}'") for d in (days or F["days"]))

def unmerge_region(ws, min_row, max_row, min_col, max_col):
    """Unmerge every merged range intersecting the region; return their bounds."""
    removed = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= max_row and rng.max_row >= min_row and \
           rng.min_col <= max_col and rng.max_col >= min_col:
            removed.append((rng.min_col, rng.min_row, rng.max_col, rng.max_row))
            ws.unmerge_cells(str(rng))
    return removed

def style_from(ws, src, dst, value=None, numfmt=None):
    c = ws[dst]
    c._style = copy(ws[src]._style)
    if value is not None:
        c.value = value
    if numfmt:
        c.number_format = numfmt
    return c

# ============================================================ 1. load + rename + global rewrite
def load_and_rewrite():
    wb = openpyxl.load_workbook(CONFIG["SOURCE_FILE"], data_only=False)
    for d in F["days"]:
        wb[f"{DONOR['ABBR']} {d}"].title = f"{A} {d}"

    rep = [
        (DONOR["NAME"].upper(), CONFIG["MONTH"].upper()),      # JULY -> AUGUST
        (DONOR["NAME"], CONFIG["MONTH"]),                      # July -> August
        (f"{DONOR['ABBR'].upper()} ", f"{A.upper()} "),        # 'JUL 1' labels
        (f"{DONOR['ABBR']} ", f"{A} "),                        # 'Jul 1' refs/labels
        (f"{CONFIG['YEAR']},{DONOR['MONTH_NUM']},", f"{CONFIG['YEAR']},{CONFIG['MONTH_NUM']},"),  # DATE()
        ("$BR$4:$BR$30", "$BR$4:$BR$29"),                      # 26 selling-day serials
    ]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    for a, b in rep:
                        if a in v:
                            v = v.replace(a, b)
                    if v != c.value:
                        c.value = v
                elif isinstance(v, dt.datetime) and v.year == CONFIG["YEAR"] and v.month == DONOR["MONTH_NUM"]:
                    c.value = v.replace(month=CONFIG["MONTH_NUM"])
    return wb

# ============================================================ 2. day tabs
def build_day_tab(wb, d):
    ws = wb[f"{A} {d}"]
    up, closed = A.upper(), d in F["closed"]

    # identity row + Sunday treatment (normalized from config, not inherited)
    ws["A1"] = f"{up} {d}"
    ws["D1"] = f"GALLATIN FORD — {up} {d} · {F['dow'][d].upper()}"
    ws["D2"] = f"Daily Deal Log · {CONFIG['MONTH']} {CONFIG['YEAR']}" + ("  ·  CLOSED (SUNDAY)" if closed else "")
    ws.sheet_properties.tabColor = (COLOR_SUNDAY if closed else
                                    COLOR_SATURDAY if F["dow"][d] == "Saturday" else COLOR_WEEKDAY)
    # nav (global rewrite fixed names; rebuild for wraparound correctness)
    prev_t = TABS[(d - 2) % CONFIG["DAYS"]]
    next_t = TABS[d % CONFIG["DAYS"]]
    ws["A3"] = f'=HYPERLINK("#\'{prev_t}\'!A1","◀  Prev Day")'
    ws["K3"] = f'=HYPERLINK("#\'{next_t}\'!A1","Next Day  ▶")'

    # glance strip rows 4-5 (SPEC §5)
    for col, lab in zip("CEGIKMO", ["UNITS", "FRONT", "FINANCE", "TOTAL", "TRADES", "PENDING", "WATER"]):
        style_from(ws, "C45", f"{col}4", lab)
    for col, ref, fmt_src in [("C", "=$G$46", "G46"), ("E", "=$G$47", "G47"), ("G", "=$G$48", "G48"),
                              ("I", "=$G$49", "G49"), ("K", "=$G$50", "G50"), ("M", "=$C$56", "C56"),
                              ("O", "=$G$57", "C56")]:
        style_from(ws, fmt_src, f"{col}5", ref)

    # AQ variance column (SPEC §5)
    style_from(ws, "AP7", "AQ7", "x ACV−MMR")
    for r in range(8, 43):
        c = style_from(ws, f"AP{r}", f"AQ{r}", f'=IF(OR($AB{r}="",$AC{r}=""),"",$AB{r}-$AC{r})',
                       numfmt="$#,##0;[Red]($#,##0)")
        c.protection = Protection(locked=True)
    ws.column_dimensions["AQ"].width = 10
    ws.column_dimensions["AQ"].hidden = False

    # FLAGS additions E54:E57 / G54:G57
    for r, lab, f_ in [
        (54, "Wholesaled", '=COUNTIFS($T$8:$T$42,"<>",$AA$8:$AA$42,"Wholesale",$AI$8:$AI$42,1)'
                           '+COUNTIFS($T$8:$T$42,"<>",$AA$8:$AA$42,"Auction",$AI$8:$AI$42,1)'),
        (55, "ACV−MMR Σ", '=SUMIFS($AQ$8:$AQ$42,$T$8:$T$42,"<>",$AI$8:$AI$42,1)'),
        (56, "ACV−MMR n", '=COUNTIFS($AB$8:$AB$42,"<>",$AC$8:$AC$42,"<>",$T$8:$T$42,"<>",$AI$8:$AI$42,1)'),
        (57, "Water", '=SUMPRODUCT(($A$8:$A$42="Used")*$AI$8:$AI$42*($Q$8:$Q$42<0))'),
    ]:
        style_from(ws, "A54", f"E{r}", lab)
        style_from(ws, "C54" if r != 55 else "C55", f"G{r}", f_)

    # D5 fix: kept-count delivered-gated
    ws["C58"] = '=COUNTIFS($T$8:$T$42,"<>",$Z$8:$Z$42,"N",$AI$8:$AI$42,1)'

    # pace block: C68 literal, C69 numerator/denominator (day DAYS keeps donor GOAL-HIT variant)
    ws["C68"] = F["left"][d]
    if d != CONFIG["DAYS"]:
        e = F["elapsed"][d]
        ws["C69"] = (f'=IF($C$64="","-",IF($C$62>={e}/GoalSellingDays*$C$64,"● AHEAD",'
                     f'IF($C$62>=0.9*$C$64*{e}/GoalSellingDays,"● ON PACE","● BEHIND")))')

    # D2 fix (defect): lock the M formula column
    for r in range(8, 43):
        ws[f"M{r}"].protection = Protection(locked=True)

    # new CF rules (SPEC §5)
    red = dict(fill=PatternFill(bgColor="C00000"), font=Font(color="FFFFFF", bold=True))
    amber = dict(fill=PatternFill(bgColor="FFD966"))
    ws.conditional_formatting.add("Q8:Q42", FormulaRule(formula=['AND($A8="Used",$AI8=1,N($Q8)<0)'], **red))
    ws.conditional_formatting.add("P8:P42", FormulaRule(formula=['AND($AI8=1,$P8="")'], **amber))
    ws.conditional_formatting.add("Q8:Q42", FormulaRule(formula=['AND($AI8=1,$Q8="")'], **amber))
    ws.conditional_formatting.add("AF8:AF42", FormulaRule(formula=['AND($AI8=1,$AF8="")'], **amber))
    ws.conditional_formatting.add("AQ8:AQ42", FormulaRule(
        formula=['AND($AQ8<>"",$AQ8>BM_ACV_OVER_MMR)'],
        fill=PatternFill(bgColor="F8CBAD"), font=Font(color="9C0006", bold=True)))

# ============================================================ 3. Staff & Lists + Benchmarks
BENCHMARKS = [  # label, defined name, value, number format
    ("VSC penetration target", "BM_VSC_PEN", 0.50, "0%"),
    ("Appraisal-to-trade floor", "BM_APPR_TRADE", 0.50, "0%"),
    ("Immediate-wholesale ceiling", "BM_IMMED_WHOLESALE", 0.33, "0%"),
    ("Used:New target", "BM_USED_TO_NEW", 1.00, "0.00"),
    ("ACV over MMR flag ($)", "BM_ACV_OVER_MMR", 500, "$#,##0"),
    ("Trade turn-in days", "BM_TURNIN_DAYS", 3, "0"),
]

def build_staff_lists(wb):
    ws = wb["Staff & Lists"]
    # move note lines out of the SalespeopleList COUNTA column
    for r in (24, 25):
        ws[f"C{r}"].value = ws[f"A{r}"].value
        ws[f"C{r}"]._style = copy(ws[f"A{r}"]._style)
        ws[f"A{r}"] = None
        ws[f"A{r}"]._style = copy(ws["A21"]._style)
    # Benchmarks block E24:F30
    style_from(ws, "A4", "E24", "BENCHMARKS — tune targets here (yellow cells)")
    for i, (lab, name, val, fmt) in enumerate(BENCHMARKS):
        r = 25 + i
        style_from(ws, "A5", f"E{r}", lab)
        c = ws[f"F{r}"]
        c.value = val
        c.number_format = fmt
        c.fill = PatternFill("solid", start_color=YELLOW, end_color=YELLOW)
        c.font = Font(name="Arial", sz=10, bold=True)
        c.protection = Protection(locked=False)
        if name in wb.defined_names:
            del wb.defined_names[name]
        wb.defined_names.add(DefinedName(name, attr_text=f"'Staff & Lists'!$F${r}"))

# ============================================================ 4. Goals
def build_goals(wb):
    ws = wb["Goals"]
    ws["C5"], ws["C6"], ws["C7"], ws["C8"] = f"{CONFIG['MONTH']} {CONFIG['YEAR']}", CONFIG["DAYS"], len(F["closed"]), CONFIG["SELLING_DAYS"]

# ============================================================ 5. Dashboard
def build_dashboard(wb):
    ws = wb["Dashboard"]
    # D4 fix: roster references
    for i in range(16):
        ws[f"AD{4 + i}"] = f"='Staff & Lists'!A{5 + i}"

    # serials + counters
    for i, d in enumerate(F["selling"]):
        ws[f"BR{4 + i}"] = F["serial"][d]
    ws["BR30"] = None
    ws["BT4"] = '=MIN(GoalSellingDays,COUNTIF($BR$4:$BR$29,"<="&TODAY()))'
    ws["M11"] = ('=IF(GoalUnitsTotal="","—",IF(BT4=0,"● MONTH NOT STARTED",'
                 'IF(BT6>=GoalUnitsTotal,"✓ GOAL HIT",'
                 'IF(BT6>=GoalUnitsTotal*BT4/GoalSellingDays,"● AHEAD OF PACE",'
                 'IF(BT6>=0.9*GoalUnitsTotal*BT4/GoalSellingDays,"● ON PACE","● BEHIND PACE")))))')

    # per-day agg new column AC (New units)
    for i, d in enumerate(F["days"]):
        style_from(ws, f"U{4 + i}", f"AC{4 + i}", f"='{A} {d}'!$C$46")

    # intel chains
    ws["BT10"] = chain("{t}!$G$57")
    ws["BT11"] = "=SUM(AC4:AC34)"
    ws["BT12"] = chain("{t}!$G$54")
    ws["BT13"] = chain("{t}!$G$55")
    ws["BT14"] = chain("{t}!$G$56")

    # pace row 12
    style_from(ws, "B10", "B12", "Units Pace")
    style_from(ws, "K11", "C12", '=IF(BT4=0,"—",ROUND(BT6/BT4*GoalSellingDays,1))', numfmt="0.0")
    style_from(ws, "M11", "D12", '=IF(OR(GoalUnitsTotal="",BT4=0),"",IF(C12>=GoalUnitsTotal,'
                                 '"▲ +"&TEXT(C12-GoalUnitsTotal,"0.0"),"▼ "&TEXT(C12-GoalUnitsTotal,"0.0")))')
    style_from(ws, "F10", "F12", "Gross Pace")
    style_from(ws, "K11", "G12", '=IF(BT4=0,"—",ROUND(BT7/BT4*GoalSellingDays,0))', numfmt="$#,##0")
    style_from(ws, "M11", "I12", '=IF(OR(GoalGrossTotal="",BT4=0),"",IF(G12>=GoalGrossTotal,'
                                 '"▲ "&TEXT(G12-GoalGrossTotal,"$#,##0"),"▼ "&TEXT(G12-GoalGrossTotal,"$#,##0")))')

    # ---- capture then rebuild the visible zone below the calendar (shift +2)
    shifted_merges = unmerge_region(ws, 26, 43, 1, 18)
    donor = {}
    for r in range(26, 44):
        for col in range(1, 19):
            L = get_column_letter(col)
            c = ws[f"{L}{r}"]
            if c.value is not None or c.has_style:
                donor[(L, r)] = (c.value, copy(c._style))
    heights = {r: ws.row_dimensions[r].height for r in range(26, 44) if r in ws.row_dimensions}
    blank_style = copy(ws["A46"]._style)
    for r in range(26, 44):
        for col in range(1, 19):
            L = get_column_letter(col)
            ws[f"{L}{r}"] = None
            ws[f"{L}{r}"]._style = copy(blank_style)
    for (L, r), (v, st) in donor.items():
        c = ws[f"{L}{r + 2}"]
        c._style = st
        c.value = v
    for r, h in heights.items():
        if h:
            ws.row_dimensions[r + 2].height = h
    for (c1, r1, c2, r2) in shifted_merges:
        ws.merge_cells(start_row=r1 + 2, start_column=c1, end_row=r2 + 2, end_column=c2)
    # regenerate the row-relative Avg/Unit family at its new rows (ranks 1-16 at rows 30-45)
    for i in range(16):
        r = 30 + i
        ws[f"L{r}"] = f'=IF(N(H{r})=0,"—",J{r}/H{r})'

    # ---- calendar rebuild: 6 week pairs, rows 15-26
    chip_st = copy(ws["H15"]._style)
    strip_st = copy(ws["H16"]._style)
    closed_st = copy(ws["B18"]._style)
    cal_merges = unmerge_region(ws, 15, 26, 2, 15)
    # donor merge geometry: how wide is a chip/strip cell? (0 = unmerged)
    chip_w = max([c2 - c1 for (c1, r1, c2, r2) in cal_merges if r1 % 2 == 1] or [0])
    strip_w = max([c2 - c1 for (c1, r1, c2, r2) in cal_merges if r1 % 2 == 0] or [0])
    for r in range(15, 27):
        for col in range(2, 16):
            L = get_column_letter(col)
            ws[f"{L}{r}"] = None
            ws[f"{L}{r}"]._style = copy(blank_style)
    for d in F["days"]:
        col, wk = dow_col(CONFIG, d), week_index(CONFIG, d)
        chip_r, strip_r = 15 + 2 * wk, 16 + 2 * wk
        ci = openpyxl.utils.column_index_from_string(col)
        c = ws[f"{col}{chip_r}"]
        c.value = f'=HYPERLINK("#\'{A} {d}\'!A1","{d}")'
        c._style = copy(chip_st)
        if chip_w:
            ws.merge_cells(start_row=chip_r, start_column=ci, end_row=chip_r, end_column=ci + chip_w)
        s = ws[f"{col}{strip_r}"]
        if d in F["closed"]:
            s.value = "closed"
            s._style = copy(closed_st)
        else:
            s.value = (f"=IF('{A} {d}'!$G$46=0,\"—\",'{A} {d}'!$G$46&\"u · $\"&"
                       f"TEXT('{A} {d}'!$G$49/1000,\"0.0\")&\"k\")")
            s._style = copy(strip_st)
        if strip_w:
            ws.merge_cells(start_row=strip_r, start_column=ci, end_row=strip_r, end_column=ci + strip_w)

    # ---- intel band rows 47-56
    lab, val = "B10", "K11"
    style_from(ws, "A26" if False else "A5", "A47", "USED DESK & F&I INTEL — MTD")
    style_from(ws, lab, "B48", "Used:New")
    style_from(ws, val, "C48", '=IF(BT11=0,"—",ROUND((BT6-BT11)/BT11,2))', numfmt="0.00")
    style_from(ws, val, "D48", '=IF(C48="—","",IF(C48>=BM_USED_TO_NEW,"▲ at/above "&TEXT(BM_USED_TO_NEW,"0.00"),'
                               '"▼ below "&TEXT(BM_USED_TO_NEW,"0.00")))')
    style_from(ws, lab, "F48", "Water deals")
    style_from(ws, val, "G48", "=BT10")
    style_from(ws, lab, "I48", "Wholesaled %")
    style_from(ws, val, "J48", '=IF(SUM(W4:W34)=0,"—",BT12/SUM(W4:W34))', numfmt="0%")
    style_from(ws, val, "K48", '=IF(J48="—","",IF(J48<=BM_IMMED_WHOLESALE,"✓ under cap",'
                               '"⚠ over "&TEXT(BM_IMMED_WHOLESALE,"0%")))')
    style_from(ws, lab, "B50", "VSC pen")
    style_from(ws, val, "C50", '=IF(BT6=0,"—",SUM(Z4:Z34)/BT6)', numfmt="0%")
    style_from(ws, val, "D50", '=IF(C50="—","",IF(C50>=BM_VSC_PEN,"✓","⚠ vs "&TEXT(BM_VSC_PEN,"0%")))')
    style_from(ws, lab, "F50", "LoJack")
    style_from(ws, val, "G50", '=IF(BT6=0,"—",SUM(AA4:AA34)/BT6)', numfmt="0%")
    style_from(ws, lab, "I50", "Int/Ext")
    style_from(ws, val, "J50", '=IF(BT6=0,"—",SUM(AB4:AB34)/BT6)', numfmt="0%")
    style_from(ws, lab, "B52", "PVR front")
    style_from(ws, val, "C52", '=IF(BT6=0,"—",SUM(U4:U34)/BT6)', numfmt="$#,##0")
    style_from(ws, lab, "E52", "back")
    style_from(ws, val, "F52", '=IF(BT6=0,"—",SUM(V4:V34)/BT6)', numfmt="$#,##0")
    style_from(ws, lab, "H52", "total")
    style_from(ws, val, "I52", '=IF(BT6=0,"—",BT7/BT6)', numfmt="$#,##0")
    style_from(ws, "A5", "A54", "F&I MANAGER SPLIT — deals · VSC · LoJack · Int/Ext")
    for i in (0, 1):
        r = 55 + i
        style_from(ws, lab, f"B{r}", f"='Staff & Lists'!G{5 + i}")
        style_from(ws, val, f"D{r}", _fni_deals_chain(r))
        for colL, prod in [("E", "AM"), ("F", "AN"), ("G", "AO")]:
            terms = "+".join(
                f"SUMPRODUCT(('{A} {d}'!$I$8:$I$42=$B{r})*'{A} {d}'!$AI$8:$AI$42*('{A} {d}'!${prod}$8:${prod}$42=\"Y\"))"
                for d in F["days"])
            style_from(ws, val, f"{colL}{r}", f'=IF($D{r}=0,"—",({terms})/$D{r})', numfmt="0%")
    # water CF
    ws.conditional_formatting.add("G48", FormulaRule(formula=["N($G$48)>0"],
                                                     fill=PatternFill(bgColor="C00000"),
                                                     font=Font(color="FFFFFF", bold=True)))

def _fni_deals_chain(r):
    terms = "+".join(f"SUMPRODUCT(('{A} {d}'!$I$8:$I$42=$B{r})*'{A} {d}'!$AI$8:$AI$42)" for d in F["days"])
    return f"={terms}"

# ============================================================ 6. Commissions
def build_commissions(wb):
    ws = wb["Commissions"]
    for i in range(16):
        ws[f"T{4 + i}"] = f"='Staff & Lists'!A{5 + i}"
    for i in range(16):
        rep = 4 + i
        for j, d in enumerate(F["days"]):
            L = get_column_letter(21 + j)          # U..AY
            t = f"'{A} {d}'"
            credit = (f"({t}!$F$8:$F$42=$T{rep})*{t}!$AJ$8:$AJ$42"
                      f"+({t}!$G$8:$G$42=$T{rep})*{t}!$AK$8:$AK$42")
            gross = f"({t}!$Q$8:$Q$42+{t}!$R$8:$R$42)"
            ws[f"{L}{rep}"] = f"=SUMPRODUCT({credit})"
            ws[f"{L}{rep + 19}"] = f"=SUMPRODUCT(({credit})*{gross})"
            ws[f"{L}{rep + 38}"] = (f"=SUMPRODUCT(({credit})*LOOKUP(({gross}*({gross}>0)),"
                                    f"'Pay Plan'!$D$5:$K$5,INDEX('Pay Plan'!$D$6:$K$12,$BB{rep},0)))")
        ws[f"BC{rep}"] = "=" + "+".join(
            f"SUMPRODUCT(('{A} {d}'!$F$8:$F$42=$T{rep})*'{A} {d}'!$AI$8:$AI$42*"
            f"(('{A} {d}'!$AM$8:$AM$42=\"Y\")+('{A} {d}'!$AN$8:$AN$42=\"Y\")+('{A} {d}'!$AO$8:$AO$42=\"Y\")))"
            for d in F["days"])
    ws["BH4"] = chain("{t}!$G$46")
    ws["BH5"] = "=" + "+".join(
        f"SUMPRODUCT('{A} {d}'!$AI$8:$AI$42*"
        f"(('{A} {d}'!$AM$8:$AM$42=\"Y\")+('{A} {d}'!$AN$8:$AN$42=\"Y\")+('{A} {d}'!$AO$8:$AO$42=\"Y\")))"
        for d in F["days"])

# ============================================================ 7. Scoreboard
def build_scoreboard(wb):
    ws = wb["Scoreboard"]
    for r in (8, 9, 10, 14, 15, 16):
        c = style_from(ws, f"E{r}", f"G{r}",
                       f'=IF(OR($C{r}="-",N($C{r})=0),"",REPT("█",MIN(20,ROUND($D{r}/$C{r}*20,0)))'
                       f'&IF($D{r}>=$C{r}," ✔",""))')
        c.number_format = "General"
        c.font = Font(name="Arial", sz=9, color="C9A227")
    style_from(ws, "A18", "B24", "PACE — projected finish")
    style_from(ws, "B19", "B25", "Units")
    style_from(ws, "D19", "C25", f'=IF(Dashboard!$BT$4=0,"-",ROUND(\'{LAST}\'!$C$62/Dashboard!$BT$4*GoalSellingDays,1))',
               numfmt="0.0")
    style_from(ws, "D19", "D25", '=IF(OR(GoalUnitsTotal="",C25="-"),"-",IF(C25>=GoalUnitsTotal,'
                                 '"▲ +"&TEXT(C25-GoalUnitsTotal,"0.0"),"▼ "&TEXT(C25-GoalUnitsTotal,"0.0")))')
    style_from(ws, "B19", "B26", "Gross")
    style_from(ws, "D19", "C26", f'=IF(Dashboard!$BT$4=0,"-",ROUND(\'{LAST}\'!$C$63/Dashboard!$BT$4*GoalSellingDays,0))',
               numfmt="$#,##0")
    style_from(ws, "D19", "D26", '=IF(OR(GoalGrossTotal="",C26="-"),"-",IF(C26>=GoalGrossTotal,'
                                 '"▲ "&TEXT(C26-GoalGrossTotal,"$#,##0"),"▼ "&TEXT(C26-GoalGrossTotal,"$#,##0")))')

# ============================================================ 8. Nightly Text
def build_nightly(wb):
    ws = wb["Nightly Text"]
    style_from(ws, "D13", "G13", "Unit pace")
    style_from(ws, "E13", "H13", '=IF($E$8=0,0,($E$12+$H$12)/$E$8*COUNT(Dashboard!$BR$4:$BR$29))', numfmt="0.0")
    style_from(ws, "E13", "F13", '=IF(OR(GoalGrossTotal="",$E$8=0),"",IF($E$13>=GoalGrossTotal,"▲","▼"))')
    style_from(ws, "E13", "I13", '=IF(OR(GoalUnitsTotal="",$E$8=0),"",IF($H$13>=GoalUnitsTotal,"▲","▼"))')
    ws["F13"].number_format = "General"
    ws["I13"].number_format = "General"

# ============================================================ 9. Trade Report
def build_trade_report(wb):
    ws = wb["Trade Report"]
    for col, lab in [("I", "Wholesaled"), ("J", "ΣACV−MMR"), ("K", "n"), ("L", "Aged")]:
        style_from(ws, "H8", f"{col}8", lab)
    for i, d in enumerate(F["days"]):
        r = 9 + i
        style_from(ws, f"C{r}", f"I{r}", f"='{A} {d}'!$G$54")
        style_from(ws, f"F{r}", f"J{r}", f"='{A} {d}'!$G$55")
        style_from(ws, f"C{r}", f"K{r}", f"='{A} {d}'!$G$56")
        style_from(ws, f"H{r}", f"L{r}", f'=IF(AND(E{r}>0,B{r}<=TODAY()-BM_TURNIN_DAYS),"⏰ AGED","")')
    for col, lab in [("H", "Capture %"), ("I", "Whlsl %"), ("J", "Avg ACV−MMR"), ("K", "Aged kept")]:
        style_from(ws, "B6", f"{col}6", lab)
    style_from(ws, "E7", "H7", '=IF(Dashboard!$BT$6=0,"—",B7/Dashboard!$BT$6)', numfmt="0%")
    style_from(ws, "E7", "I7", '=IF(B7=0,"—",SUM(I9:I39)/B7)', numfmt="0%")
    style_from(ws, "G7", "J7", '=IF(SUM(K9:K39)=0,"—",SUM(J9:J39)/SUM(K9:K39))', numfmt="$#,##0")
    style_from(ws, "B7", "K7", '=COUNTIF(L9:L39,"⏰ AGED")')
    ws.column_dimensions["K"].width = 5
    for c in ("I", "J", "L"):
        ws.column_dimensions[c].width = 11
    # recap verdict CF
    ws.conditional_formatting.add("H7", FormulaRule(formula=['AND(H7<>"—",N(H7)<BM_APPR_TRADE)'],
                                                    fill=PatternFill(bgColor="FFD966")))
    ws.conditional_formatting.add("I7", FormulaRule(formula=['AND(I7<>"—",N(I7)>BM_IMMED_WHOLESALE)'],
                                                    fill=PatternFill(bgColor="FFD966")))
    ws.conditional_formatting.add("J7", FormulaRule(formula=['AND(J7<>"—",N(J7)>BM_ACV_OVER_MMR)'],
                                                    fill=PatternFill(bgColor="F8CBAD")))

# ============================================================ 10. Deal Explorer
def build_explorer(wb):
    ws = wb["Deal Explorer"]
    ws.freeze_panes = "A8"                      # D3 fix
    # EX fix (flagged at G3): July's ✓ gate `IF(C8="",...)` is engine-dependent for empty
    # rows (Excel: link to empty yields 0 -> every empty row shows ✓; LibreOffice: yields
    # empty). Gate on both, and count booked deals instead of COUNTA (which counts the 0s).
    for r in range(8, 1093):
        c = ws[f"A{r}"]
        if isinstance(c.value, str) and c.value.startswith(f'=IF(C{r}="",'):
            c.value = c.value.replace(f'=IF(C{r}="",', f'=IF(OR(C{r}="",C{r}=0),', 1)
    ws["B6"] = ('="Showing "&COUNTIF(A8:A1092,"✓")&" of "&'
                'SUMPRODUCT(($C$8:$C$1092<>"")*($C$8:$C$1092<>0))&" deals booked this month"')
    dv = DataValidation(type="list", formula1="SalespeopleList", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B5")
    ws.auto_filter.ref = None
    t = Table(displayName="DealLedger", ref="A7:L1092")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws.add_table(t)

# ============================================================ 11. Guide
GUIDE_ROWS = [
    ("Daily flow", "Open today’s day tab (gold TODAY chip, or click a calendar day on the Dashboard). One deal per "
     "row: type, stock, deal#, customer, SP1/SP2, vehicle, gross, lender, ad source, status. Splits: put both reps "
     "in SP1 and SP2 — credit auto-splits ½/½. The strip at the top of every day tab shows today at a glance."),
    ("Status drives everything", "Any deal with a Deal# counts (units, gross, pace, pay) unless its Status is "
     "Pending, Unwind or Back Out. That one rule now runs the whole book — Commissions included."),
    ("Flags on the day tab", "Red Front Gross on a used deal = water (losing money on the front). Amber Lender / "
     "Front Gross / Ad Source = a delivered deal missing that field. Red ACV−MMR = we allowed more than the "
     "over-allowance limit above MMR (tune limits on Staff & Lists → BENCHMARKS)."),
    ("Dashboard", "Live month roll-up: MTD cards, goal & pace band with projected finish (▲/▼ vs goal), clickable "
     "calendar, salesperson leaderboard, channel mix, products, and the USED DESK & F&I INTEL band — used:new "
     "ratio, water count, wholesale %, product penetration vs target, PVR, per-F&I-manager split."),
    ("Deal Explorer", "Every deal in one table. Use the header filter arrows, or set the criteria row (Salesperson "
     "matches SP1 or SP2, Type, Status, Certified, Min Gross, Search) and filter the ✓ column."),
    ("Commissions", "Retro pay: final monthly units set the tier for every deal. Draw is entered by payroll "
     "(yellow column); Balance = Earnings − Draw. Six internal audits sit under the table."),
    ("Trade Report", "Per-day trades, turned-in vs kept, ACV, wholesaled count, ACV−MMR spread and aged-kept "
     "flags (⏰ after the turn-in deadline). Recap row grades appraisal capture and immediate-wholesale % "
     "against the benchmarks."),
    ("Nightly Text", "Fill the yellow live inputs, wait for “✅ READY TO COPY”, then copy the box at the bottom "
     "into Teams. Day-override cell re-runs a prior day. Pace cells show projected units and gross."),
    ("Goals, Benchmarks & Pay Plan", "Set monthly goals on Goals (yellow). Tune targets on Staff & Lists → "
     "BENCHMARKS (yellow) — every flag and verdict reads them live. Pay Plan (hidden tab) is the single source "
     "of truth for commission math."),
]

def build_guide(wb):
    ws = wb["Guide"]
    for i, (b, c) in enumerate(GUIDE_ROWS):
        r = 7 + i
        style_from(ws, "B7", f"B{r}", b)
        style_from(ws, "C7", f"C{r}", c)

# ============================================================ 12. test matrix (Gate 3)
TEST_DEALS = [  # column -> value maps, row 8..14 of day-1 tab
    dict(A="New", B="N1001", C=12, D=5001, E="ALPHA", F="FEDY ALI", H="NATA RODRIGUEZ", I="JOHN MCCAN",
         J=2026, K="FORD", L="F-150", N="N", O="Y", P="Ford Credit", Q=2000, R=1500,
         T="T1001", U=2022, V="FORD", W="Escape", X=45000, Y="1FMCU9G67NUA12345", Z="Y", AA="Retail",
         AB=15000, AC=14200, AD=14800, AE=14500, AF="GOOGLE", AG="Delivered", AH="N", AM="Y"),
    dict(A="Used", B="U2002", D=5002, E="BRAVO", F="CHAD EMBRY", G="RICH AMMERMAN", H="TONY BATTISTA",
         I="JACOB MAGED", J=2023, K="TOYOTA", L="Camry", N="N", P="CapitalOne", Q=1200, R=800,
         AF="BE BACK", AG="Delivered"),
    dict(A="Used", B="U2003", D=5003, E="CHARLIE", F="ZAY JOHNSON", H="JOE BARBA", I="JOHN MCCAN",
         J=2021, K="HONDA", L="Civic", N="N", P="Chase", Q=-500, R=900, AF="CAR GURU", AG="Delivered"),
    dict(A="New", B="N1004", D=5004, E="DELTA", F="FEDY ALI", H="NATA RODRIGUEZ", J=2026, K="FORD",
         L="Bronco", N="N", Q=3000, AF="AUTOTRADER", AG="Pending"),
    dict(A="Used", B="U2005", D=5005, E="ECHO", F="CHAD EMBRY", H="TONY BATTISTA", J=2020, K="FORD",
         L="Fusion", N="N", Q=1000, R=500, AF="REFERRAL", AG="Unwind"),
    dict(A="New", B="N1006", D=5006, E="FOXTROT", F="MERCEDES CLARK", H="NATA RODRIGUEZ", J=2026,
         K="FORD", L="Maverick", N="N", P="CASH", Q=1800, AF="FRESH UP", AG="Delivered"),
    dict(A="Used", B="U2007", D=5007, E="GOLF", F="FEDY ALI", H="JOE BARBA", I="JACOB MAGED", J=2024,
         K="FORD", L="Explorer", N="Y", P="Ford Credit", Q=6000, R=3000, AF="GOOGLE", AG="Delivered", AM="Y"),
]

def seed_test_deals(wb):
    ws = wb[TABS[0]]
    for i, deal in enumerate(TEST_DEALS):
        r = 8 + i
        for col, v in deal.items():
            ws[f"{col}{r}"] = v

# ============================================================ main
def sanity(wb):
    """In-generator QA: no donor-month residue in any string cell."""
    bad = []
    pat = re.compile(rf"\b{DONOR['ABBR']}\b|{DONOR['NAME']}|{CONFIG['YEAR']},{DONOR['MONTH_NUM']},")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and pat.search(c.value):
                    bad.append(f"{ws.title}!{c.coordinate}: {c.value[:60]}")
    if bad:
        raise SystemExit("DONOR RESIDUE:\n" + "\n".join(bad[:20]))

def main(test_deals=False):
    wb = load_and_rewrite()
    build_staff_lists(wb)
    build_goals(wb)
    for d in F["days"]:
        build_day_tab(wb, d)
    build_dashboard(wb)
    build_commissions(wb)
    build_scoreboard(wb)
    build_nightly(wb)
    build_trade_report(wb)
    build_explorer(wb)
    build_guide(wb)
    if test_deals:
        seed_test_deals(wb)
    sanity(wb)
    wb.save(CONFIG["OUTPUT_FILE"])
    print(f"saved {CONFIG['OUTPUT_FILE'].name} (test_deals={test_deals})")

if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--test-deals", action="store_true")
    main(test_deals=p.parse_args().test_deals)
